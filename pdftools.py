"""Document-processing core for EasyPDFTools.

Everything in this module is plain synchronous Python with no web-framework imports,
so it can be unit-tested on its own. main.py calls these functions from worker threads.

Errors that are the user's fault (corrupt file, password-protected PDF, scanned PDF
with no text, ...) are raised as ToolError with a friendly message and an HTTP status.
"""
from __future__ import annotations

import io
import logging
import os
import re
import threading
import warnings
import zipfile
from datetime import date, datetime
from datetime import time as dtime
from pathlib import Path
from xml.sax.saxutils import escape

import docx
import openpyxl
import pypdf
import pypdfium2 as pdfium
from docx.oxml import OxmlElement
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from PIL import Image, ImageOps
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

try:  # Optional: proper Arabic shaping + right-to-left ordering in generated PDFs.
    import arabic_reshaper
except ImportError:  # pragma: no cover
    arabic_reshaper = None
try:
    from bidi import get_display
except ImportError:  # pragma: no cover
    try:
        from bidi.algorithm import get_display
    except ImportError:
        get_display = None

log = logging.getLogger("easypdftools")
logging.getLogger("pypdf").setLevel(logging.ERROR)


def _env_int(name: str, default: int) -> int:
    try:
        return int(os.getenv(name, str(default)))
    except ValueError:
        return default


MAX_PDF_PAGES = _env_int("MAX_PDF_PAGES", 500)
MAX_JPG_PAGES = _env_int("MAX_JPG_PAGES", 100)
MAX_EXCEL_ROWS = _env_int("MAX_EXCEL_ROWS", 5000)
OCR_MAX_PAGES = _env_int("OCR_MAX_PAGES", 2)
MAX_UNZIPPED_BYTES = 200 * 1024 * 1024  # guards against zip bombs in .docx/.xlsx

NO_TEXT_MSG = (
    "No selectable text was found. This looks like a scanned PDF, "
    "so try the OCR tool (Scan to Word) instead."
)


class ToolError(Exception):
    """A problem the user can understand and fix (bad file, limits, ...)."""

    def __init__(self, message: str, status: int = 400):
        super().__init__(message)
        self.message = message
        self.status = status


# ----------------------------------------------------------------------------
# Shared helpers
# ----------------------------------------------------------------------------

# Characters that are illegal in XML 1.0 (python-docx / openpyxl raise on them).
_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_RTL_CHARS = re.compile(r"[\u0590-\u08ff\ufb1d-\ufdff\ufe70-\ufeff]")
_LETTERS = re.compile(r"[^\W\d_]")


def clean(text: str | None) -> str:
    return _CTRL.sub("", text or "")


def has_rtl(text: str) -> bool:
    return bool(_RTL_CHARS.search(text))


def is_rtl(text: str) -> bool:
    letters = len(_LETTERS.findall(text))
    return letters > 0 and len(_RTL_CHARS.findall(text)) / letters > 0.4


def ensure_pdf(content: bytes) -> None:
    if b"%PDF-" not in content[:1024]:
        raise ToolError("This file doesn't look like a PDF.", 400)


def open_pdf(content: bytes) -> pypdf.PdfReader:
    """Open a PDF with pypdf, turning every failure into a friendly ToolError."""
    ensure_pdf(content)
    try:
        reader = pypdf.PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            try:
                unlocked = reader.decrypt("")
            except Exception:
                unlocked = 0
            if not unlocked:
                raise ToolError(
                    "This PDF is password-protected. Remove the password and try again.",
                    400,
                )
        page_count = len(reader.pages)
    except ToolError:
        raise
    except Exception as exc:
        log.info("Unreadable PDF: %s", exc)
        raise ToolError("We couldn't read this PDF. The file may be corrupted.", 400)
    if page_count == 0:
        raise ToolError("This PDF has no pages.", 400)
    if page_count > MAX_PDF_PAGES:
        raise ToolError(
            f"This PDF has {page_count} pages; the limit is {MAX_PDF_PAGES} per file.", 413
        )
    return reader


_OLE_MAGIC = bytes.fromhex("D0CF11E0A1B11AE1")


def check_zip(content: bytes, kind: str) -> None:
    if content.startswith(_OLE_MAGIC):
        # Old .xls/.doc files and password-protected .xlsx/.docx files share this container.
        if "EncryptedPackage".encode("utf-16-le") in content:
            raise ToolError(
                "This file is password-protected. In Office, open it, choose File > Info > "
                "Protect > Encrypt with Password, clear the password, save, and try again.",
                400,
            )
        raise ToolError(
            f"This is an older-format Office file. Open it, use Save As to convert it to {kind}, "
            "and try again.",
            400,
        )
    if not content.startswith(b"PK\x03\x04"):
        raise ToolError(f"This doesn't look like a valid {kind} file.", 400)
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as zf:
            total = sum(info.file_size for info in zf.infolist())
    except zipfile.BadZipFile:
        raise ToolError(f"This {kind} file appears to be corrupted.", 400)
    if total > MAX_UNZIPPED_BYTES:
        raise ToolError("This file is too large once unpacked.", 413)


def _page_text(page: pypdf.PageObject, layout: bool = False) -> str:
    try:
        text = page.extract_text(extraction_mode="layout") if layout else page.extract_text()
    except Exception as exc:
        log.warning("Text extraction failed on a page: %s", exc)
        return ""
    return clean(text)


def _save_docx(document) -> bytes:
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _set_rtl(paragraph) -> None:
    """Mark a python-docx paragraph as right-to-left."""
    try:
        p_pr = paragraph._p.get_or_add_pPr()
        p_pr.append(OxmlElement("w:bidi"))
    except Exception:  # pragma: no cover - cosmetic only
        pass


# ----------------------------------------------------------------------------
# PDF -> Word
# ----------------------------------------------------------------------------

_COL_SPLIT = re.compile(r"\t+|\s{3,}")
_BULLET = re.compile(r"^(?:[\u2022\u25cf\u25aa\u25e6\u2023\u00b7*\-\u2013]|\d{1,3}[.)])\s+")
_HEADING_END = (".", ",", ";", ":", "!", "?")


def _split_cols(line: str) -> list[str]:
    return [c.strip() for c in _COL_SPLIT.split(line) if c.strip()]


def _table_run(block: list[str], start: int) -> list[list[str]] | None:
    """Consecutive lines with the same 2-8 column layout form a table."""
    first = _split_cols(block[start])
    if not 2 <= len(first) <= 8:
        return None
    rows = [first]
    for line in block[start + 1 :]:
        cols = _split_cols(line)
        if len(cols) != len(first):
            break
        rows.append(cols)
    return rows if len(rows) >= 2 else None


def _join_lines(prev: str, nxt: str) -> str:
    if prev.endswith("-") and len(prev) > 1 and prev[-2].isalpha() and nxt[:1].islower():
        return prev[:-1] + nxt  # de-hyphenate a word split across lines
    return f"{prev} {nxt}"


def _looks_like_heading(text: str) -> bool:
    return (
        len(text) <= 80
        and len(text.split()) <= 12
        and not text.endswith(_HEADING_END)
        and not text[:1].islower()
        and not _BULLET.match(text)
    )


def _write_text(document, lines: list[str], width: int) -> None:
    items: list[list] = []  # [style, text, length of last physical line]
    for line in lines:
        m = _BULLET.match(line)
        if m:
            style = "List Number" if m.group(0).strip()[0].isdigit() else "List Bullet"
            items.append([style, line[m.end() :].strip(), len(line)])
        elif items and items[-1][2] >= 0.75 * width:  # previous line was full width: it wrapped
            items[-1][1] = _join_lines(items[-1][1], line)
            items[-1][2] = len(line)
        else:
            items.append([None, line, len(line)])

    single = len(lines) == 1
    for style, text, _ in items:
        if style:
            para = document.add_paragraph(text, style=style)
        elif single and _looks_like_heading(text):
            para = document.add_heading(text, level=1 if text.isupper() else 2)
        else:
            para = document.add_paragraph(text)
        if is_rtl(text):
            _set_rtl(para)


def _write_page(document, text: str) -> None:
    lines = [ln.strip() for ln in text.split("\n")]
    lengths = sorted(len(ln) for ln in lines if ln)
    width = lengths[int(0.9 * (len(lengths) - 1))] if lengths else 0

    blocks: list[list[str]] = []
    current: list[str] = []
    for ln in lines:
        if ln:
            current.append(ln)
        elif current:
            blocks.append(current)
            current = []
    if current:
        blocks.append(current)

    for block in blocks:
        pending: list[str] = []
        i = 0
        while i < len(block):
            rows = _table_run(block, i)
            if rows:
                if pending:
                    _write_text(document, pending, width)
                    pending = []
                table = document.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Table Grid"
                for r, row in enumerate(rows):
                    for c, cell_text in enumerate(row):
                        table.cell(r, c).text = cell_text
                document.add_paragraph("")
                i += len(rows)
            else:
                pending.append(block[i])
                i += 1
        if pending:
            _write_text(document, pending, width)


def pdf_to_docx(content: bytes) -> bytes:
    reader = open_pdf(content)
    document = docx.Document()
    wrote_any = False
    for page in reader.pages:
        text = _page_text(page, layout=True)
        if not text.strip():
            continue
        if wrote_any:
            document.add_page_break()
        wrote_any = True
        _write_page(document, text)
    if not wrote_any:
        raise ToolError(NO_TEXT_MSG, 422)
    return _save_docx(document)


# ----------------------------------------------------------------------------
# PDF -> Excel
# ----------------------------------------------------------------------------

_XL_SPLIT = re.compile(r"\t+|\s{2,}")
_NUMBER = re.compile(r"^-?\d{1,3}(?:,\d{3})+(?:\.\d+)?$|^-?\d+(?:\.\d+)?$")


def _coerce_number(text: str):
    """Turn '1,234.50' into a number, but keep IDs/phone numbers/leading zeros as text."""
    s = text.strip()
    if not _NUMBER.match(s):
        return text
    unsigned = s.lstrip("-")
    if len(re.sub(r"\D", "", s)) > 15:
        return text
    if len(unsigned) > 1 and unsigned[0] == "0" and not unsigned.startswith("0."):
        return text
    n = s.replace(",", "")
    return float(n) if "." in n else int(n)


def _set_cell(cell, text: str) -> None:
    cell.value = _coerce_number(text)
    if cell.data_type == "f":  # text starting with "=" must never become a formula
        cell.data_type = "s"


def pdf_to_xlsx(content: bytes) -> bytes:
    reader = open_pdf(content)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Extracted"
    bold = Font(bold=True)
    widths: dict[int, int] = {}
    row = 1
    wrote_any = False
    for number, page in enumerate(reader.pages, start=1):
        text = _page_text(page, layout=True)
        if not text.strip():
            continue
        wrote_any = True
        header = ws.cell(row=row, column=1, value=f"Page {number}")
        header.font = bold
        row += 1
        for line in text.split("\n"):
            line = line.strip()
            if not line:
                continue
            parts = [p.strip() for p in _XL_SPLIT.split(line) if p.strip()]
            for col, part in enumerate(parts, start=1):
                _set_cell(ws.cell(row=row, column=col), part)
                widths[col] = max(widths.get(col, 0), len(part))
            row += 1
        row += 1
    if not wrote_any:
        raise ToolError(NO_TEXT_MSG, 422)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 8), 60)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ----------------------------------------------------------------------------
# PDF -> text
# ----------------------------------------------------------------------------


def pdf_to_txt(content: bytes) -> bytes:
    reader = open_pdf(content)
    text = "\n\n".join(_page_text(p) for p in reader.pages).strip()
    if not text:
        raise ToolError(NO_TEXT_MSG, 422)
    return text.encode("utf-8")


# ----------------------------------------------------------------------------
# PDF rendering (PDF -> JPG and the OCR tool). pdfium is not thread-safe, so every
# pdfium call is serialised behind one lock.
# ----------------------------------------------------------------------------

_PDFIUM_LOCK = threading.Lock()


def _open_pdfium(content: bytes):
    ensure_pdf(content)
    with _PDFIUM_LOCK:
        try:
            return pdfium.PdfDocument(content)
        except Exception as exc:
            if "password" in str(exc).lower():
                raise ToolError(
                    "This PDF is password-protected. Remove the password and try again.", 400
                )
            log.info("pdfium could not open file: %s", exc)
            raise ToolError("We couldn't read this PDF. The file may be corrupted.", 400)


def _close_pdfium(pdf) -> None:
    with _PDFIUM_LOCK:
        try:
            pdf.close()
        except Exception:  # pragma: no cover
            pass


def _render_page(pdf, index: int, dpi: int, max_px: int = 4000) -> Image.Image:
    with _PDFIUM_LOCK:
        page = pdf[index]
        try:
            width, height = page.get_size()
            scale = min(dpi / 72, max_px / max(width, height, 1))
            return page.render(scale=scale).to_pil().convert("RGB")
        finally:
            page.close()


def pdf_to_jpg_zip(content: bytes, dpi: int = 150, quality: int = 85) -> bytes:
    pdf = _open_pdfium(content)
    try:
        total = len(pdf)
        if total == 0:
            raise ToolError("This PDF has no pages.", 400)
        if total > MAX_JPG_PAGES:
            raise ToolError(
                f"This PDF has {total} pages; PDF to JPG supports up to {MAX_JPG_PAGES} "
                "pages at a time. Split the file first.",
                413,
            )
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_STORED) as zf:  # JPEGs don't compress further
            for i in range(total):
                image = _render_page(pdf, i, dpi)
                jpg = io.BytesIO()
                image.save(jpg, "JPEG", quality=quality, optimize=True)
                zf.writestr(f"page_{i + 1:03d}.jpg", jpg.getvalue())
        return buf.getvalue()
    finally:
        _close_pdfium(pdf)


def pdf_to_jpeg_pages(
    content: bytes, max_pages: int, dpi: int = 200, quality: int = 90
) -> tuple[list[bytes], int]:
    """Render the first `max_pages` pages to JPEG bytes. Returns (images, total_pages)."""
    pdf = _open_pdfium(content)
    try:
        total = len(pdf)
        if total == 0:
            raise ToolError("This PDF has no pages.", 400)
        images = []
        for i in range(min(total, max_pages)):
            buf = io.BytesIO()
            _render_page(pdf, i, dpi).save(buf, "JPEG", quality=quality)
            images.append(buf.getvalue())
        return images, total
    finally:
        _close_pdfium(pdf)


# ----------------------------------------------------------------------------
# OCR (Google Vision response parsing + Word output)
# ----------------------------------------------------------------------------


def vision_paragraphs(annotation: dict) -> list[str]:
    """Rebuild paragraphs from a Vision `fullTextAnnotation` (falls back to plain text)."""
    paragraphs: list[str] = []
    for page in annotation.get("pages", []):
        for block in page.get("blocks", []):
            for para in block.get("paragraphs", []):
                buf: list[str] = []
                for word in para.get("words", []):
                    for symbol in word.get("symbols", []):
                        buf.append(symbol.get("text", ""))
                        kind = symbol.get("property", {}).get("detectedBreak", {}).get("type")
                        if kind in ("SPACE", "SURE_SPACE", "EOL_SURE_SPACE", "LINE_BREAK"):
                            buf.append(" ")
                text = clean("".join(buf)).strip()
                if text:
                    paragraphs.append(text)
    if not paragraphs and annotation.get("text"):
        paragraphs = [p.strip() for p in clean(annotation["text"]).split("\n") if p.strip()]
    return paragraphs


def build_ocr_docx(pages: list[list[str]]) -> bytes:
    document = docx.Document()
    wrote_any = False
    for paragraphs in pages:
        texts = [t for t in (clean(p).strip() for p in paragraphs) if t]
        if not texts:
            continue
        if wrote_any:
            document.add_page_break()
        wrote_any = True
        for text in texts:
            para = document.add_paragraph(text)
            if is_rtl(text):
                _set_rtl(para)
    if not wrote_any:
        raise ToolError("No text was detected in the pages we scanned.", 422)
    return _save_docx(document)


# ----------------------------------------------------------------------------
# Fonts + right-to-left helpers for generated PDFs
# ----------------------------------------------------------------------------

_FONT_DIRS = (
    "/usr/share/fonts/truetype/dejavu",
    "/usr/share/fonts/dejavu",
    "/usr/share/fonts/TTF",
)


def _register_fonts() -> tuple[str, str]:
    regular = os.getenv("PDF_FONT_REGULAR")
    bold = os.getenv("PDF_FONT_BOLD")
    if not regular:
        for directory in _FONT_DIRS:
            candidate = Path(directory) / "DejaVuSans.ttf"
            if candidate.exists():
                regular = str(candidate)
                bold_candidate = Path(directory) / "DejaVuSans-Bold.ttf"
                bold = bold or (str(bold_candidate) if bold_candidate.exists() else None)
                break
    if not regular or not Path(regular).exists():
        log.warning("No Unicode TTF font found; PDFs fall back to Helvetica (Latin only).")
        return "Helvetica", "Helvetica-Bold"
    try:
        pdfmetrics.registerFont(TTFont("AppSans", regular))
        bold_name = "AppSans"
        if bold and Path(bold).exists():
            pdfmetrics.registerFont(TTFont("AppSans-Bold", bold))
            bold_name = "AppSans-Bold"
        pdfmetrics.registerFontFamily(
            "AppSans", normal="AppSans", bold=bold_name, italic="AppSans", boldItalic=bold_name
        )
        return "AppSans", bold_name
    except Exception:
        log.exception("Font registration failed; falling back to Helvetica.")
        return "Helvetica", "Helvetica-Bold"


BASE_FONT, BOLD_FONT = _register_fonts()


def _shape(text: str) -> str:
    """Join Arabic letters into their connected forms (logical order is kept)."""
    if arabic_reshaper and has_rtl(text):
        return arabic_reshaper.reshape(text)
    return text


def _visual(text: str) -> str:
    """Reorder bidirectional text into display order for reportlab (which draws left to right)."""
    if get_display and has_rtl(text):
        return get_display(text)
    return text


def _cell_markup(text: str) -> str:
    return escape(_visual(_shape(text))).replace("\n", "<br/>")


def _rtl_paragraph(text: str, style: ParagraphStyle, width: float) -> list:
    """Wrap on logical order first, then reorder each line, so RTL paragraphs read correctly."""
    mid = ParagraphStyle(f"{style.name}-rtl", parent=style, alignment=TA_RIGHT, spaceAfter=0)
    last = ParagraphStyle(f"{style.name}-rtl-last", parent=style, alignment=TA_RIGHT)
    font, size = style.fontName, style.fontSize
    space = pdfmetrics.stringWidth(" ", font, size)
    lines: list[list[str]] = []
    current: list[str] = []
    used = 0.0
    for word in (_shape(w) for w in text.split()):
        w = pdfmetrics.stringWidth(word, font, size)
        if current and used + space + w > width:
            lines.append(current)
            current, used = [word], w
        else:
            used += (space if current else 0) + w
            current.append(word)
    if current:
        lines.append(current)
    return [
        Paragraph(escape(_visual(" ".join(line))), last if i == len(lines) - 1 else mid)
        for i, line in enumerate(lines)
    ]


def _styles() -> dict:
    body = ParagraphStyle("body", fontName=BASE_FONT, fontSize=10.5, leading=14.5, spaceAfter=6)

    def heading(name: str, size: int) -> ParagraphStyle:
        return ParagraphStyle(
            name, parent=body, fontName=BOLD_FONT, fontSize=size, leading=size * 1.25,
            spaceBefore=8, spaceAfter=6,
        )

    return {
        "body": body,
        "title": heading("title", 22),
        "h1": heading("h1", 18),
        "h2": heading("h2", 14.5),
        "h3": heading("h3", 12),
        "list": ParagraphStyle("list", parent=body, leftIndent=18, bulletIndent=4),
        "cache": {},
    }


def _aligned(styles: dict, key: str, align: int | None) -> ParagraphStyle:
    if align is None:
        return styles[key]
    cache_key = (key, align)
    if cache_key not in styles["cache"]:
        styles["cache"][cache_key] = ParagraphStyle(
            f"{key}-{align}", parent=styles[key], alignment=align
        )
    return styles["cache"][cache_key]


# ----------------------------------------------------------------------------
# Word -> PDF
# ----------------------------------------------------------------------------

_ALIGN = {1: TA_CENTER, 2: TA_RIGHT, 3: TA_JUSTIFY}


def _docx_page(src) -> tuple[tuple[float, float], tuple[float, float, float, float]]:
    try:
        section = src.sections[0]
        width, height = section.page_width.pt, section.page_height.pt
        margins = tuple(
            m.pt
            for m in (
                section.left_margin,
                section.right_margin,
                section.top_margin,
                section.bottom_margin,
            )
        )
        if width > 200 and height > 200 and margins[0] + margins[1] < width - 150:
            return (width, height), margins  # type: ignore[return-value]
    except Exception:
        pass
    return A4, (54.0, 54.0, 54.0, 54.0)


def _runs_markup(paragraph: DocxParagraph) -> str:
    runs_text = "".join(r.text for r in paragraph.runs)
    if runs_text != paragraph.text:  # e.g. hyperlinks: python-docx keeps their text only in .text
        return escape(clean(paragraph.text)).replace("\t", "&nbsp;" * 4).replace("\n", "<br/>")
    parts = []
    for run in paragraph.runs:
        text = clean(run.text)
        if not text:
            continue
        text = escape(text).replace("\t", "&nbsp;" * 4).replace("\n", "<br/>")
        if run.bold:
            text = f"<b>{text}</b>"
        if run.italic:
            text = f"<i>{text}</i>"
        if run.underline:
            text = f"<u>{text}</u>"
        parts.append(text)
    return "".join(parts)


def _has_numbering(paragraph: DocxParagraph) -> bool:
    try:
        p_pr = paragraph._p.pPr
        return p_pr is not None and p_pr.numPr is not None
    except Exception:
        return False


def _add_docx_paragraph(story: list, para: DocxParagraph, styles: dict, width: float, state: dict):
    raw = clean(para.text)
    if not raw.strip():
        story.append(Spacer(1, 6))
        state["number"] = 0
        return

    style_name = (para.style.name if para.style is not None else "") or ""
    key = "body"
    if style_name == "Title":
        key = "title"
    else:
        m = re.match(r"Heading (\d)", style_name)
        if m:
            key = f"h{min(int(m.group(1)), 3)}"

    is_number = style_name.startswith("List Number")
    is_bullet = style_name.startswith("List Bullet") or (not is_number and _has_numbering(para))
    state["number"] = state["number"] + 1 if is_number else 0

    try:
        align = _ALIGN.get(int(para.alignment)) if para.alignment is not None else None
    except Exception:
        align = None

    if is_rtl(raw):
        story.extend(_rtl_paragraph(raw, styles[key], width - 12))
        return

    markup = _shape(_runs_markup(para)) if has_rtl(raw) else _runs_markup(para)
    if is_number or is_bullet:
        bullet = f"{state['number']}." if is_number else "\u2022"
        story.append(Paragraph(markup, _aligned(styles, "list", align), bulletText=bullet))
    else:
        story.append(Paragraph(markup, _aligned(styles, key, align)))


def _add_docx_table(story: list, table: DocxTable, styles: dict, width: float) -> None:
    rows = []
    for row in table.rows:
        rows.append([clean(cell.text) for cell in row.cells])
    if not rows or not rows[0]:
        return
    columns = max(len(r) for r in rows)
    size = 10 if columns <= 4 else 8.5
    cell_style = ParagraphStyle(
        "cell", parent=styles["body"], fontSize=size, leading=size * 1.3, spaceAfter=0
    )
    data = [
        [Paragraph(_cell_markup(c), cell_style) for c in row] + [""] * (columns - len(row))
        for row in rows
    ]
    flowable = Table(data, colWidths=[(width - 12) / columns] * columns, repeatRows=1)
    flowable.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(flowable)
    story.append(Spacer(1, 8))


def docx_to_pdf(content: bytes) -> bytes:
    check_zip(content, "Word (.docx)")
    try:
        src = docx.Document(io.BytesIO(content))
    except Exception as exc:
        log.info("Unreadable docx: %s", exc)
        raise ToolError(
            "We couldn't read this Word file. Make sure it's a .docx (not an old .doc).", 400
        )

    page_size, (left, right, top, bottom) = _docx_page(src)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=page_size, leftMargin=left, rightMargin=right,
        topMargin=top, bottomMargin=bottom,
    )
    styles = _styles()
    state = {"number": 0}
    story: list = []
    for child in src.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            _add_docx_paragraph(story, DocxParagraph(child, src), styles, doc.width, state)
        elif tag == "tbl":
            _add_docx_table(story, DocxTable(child, src), styles, doc.width)
            state["number"] = 0
    if not any(isinstance(f, (Paragraph, Table)) for f in story):
        raise ToolError("This Word document has no text to convert.", 422)
    try:
        doc.build(story)
    except Exception:
        log.exception("reportlab failed to build the PDF")
        raise ToolError("We couldn't convert this document because of an unsupported layout.", 422)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# Excel -> PDF
# ----------------------------------------------------------------------------


def _format_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else f"{value:.10g}"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d" if value.time() == dtime(0) else "%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, dtime):
        return value.strftime("%H:%M")
    return clean(str(value))[:500]


def xlsx_to_pdf(content: bytes) -> bytes:
    check_zip(content, "Excel (.xlsx)")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        log.info("Unreadable xlsx: %s", exc)
        raise ToolError(
            "We couldn't read this Excel file. Make sure it's an .xlsx workbook.", 400
        )

    sheets: list[tuple[str, list[list[str]]]] = []
    total_rows = 0
    truncated = False
    try:
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            rows: list[list[str]] = []
            for scanned, raw in enumerate(ws.iter_rows(values_only=True)):
                if scanned > 200_000:  # sheets with huge formatted-but-empty ranges
                    break
                cells = [_format_cell(c) for c in raw]
                while cells and cells[-1] == "":
                    cells.pop()
                if not cells:
                    continue
                rows.append(cells)
                total_rows += 1
                if total_rows >= MAX_EXCEL_ROWS:
                    truncated = True
                    break
            if rows:
                sheets.append((ws.title, rows))
            if truncated:
                break
    finally:
        wb.close()
    if not sheets:
        raise ToolError("This workbook has no data to convert.", 422)

    max_cols = max(len(r) for _, rows in sheets for r in rows)
    page = landscape(A4) if max_cols > 6 else A4
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=page, leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30
    )
    size = 9.5 if max_cols <= 6 else 8.5 if max_cols <= 10 else 7.5 if max_cols <= 14 else 6.5
    cell_style = ParagraphStyle(
        "xcell", fontName=BASE_FONT, fontSize=size, leading=size * 1.25
    )
    head_style = ParagraphStyle("xhead", parent=cell_style, fontName=BOLD_FONT)
    title_style = ParagraphStyle(
        "xtitle", fontName=BOLD_FONT, fontSize=12, leading=15, spaceAfter=6
    )

    story: list = []
    for index, (title, rows) in enumerate(sheets):
        if index:
            story.append(PageBreak())
        if len(sheets) > 1:
            story.append(Paragraph(_cell_markup(title), title_style))
        columns = max(len(r) for r in rows)
        weights = [4] * columns
        for r in rows[:300]:
            for c, text in enumerate(r):
                weights[c] = max(weights[c], min(len(text), 40))
        scale = (doc.width - 12) / sum(weights)
        data = [
            [
                Paragraph(_cell_markup(text), head_style if r == 0 else cell_style)
                for text in row + [""] * (columns - len(row))
            ]
            for r, row in enumerate(rows)
        ]
        table = Table(data, colWidths=[w * scale for w in weights], repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e5e7eb")),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        story.append(table)
    if truncated:
        story.append(Spacer(1, 8))
        story.append(
            Paragraph(f"Output truncated after {MAX_EXCEL_ROWS} rows.", cell_style)
        )
    try:
        doc.build(story)
    except Exception:
        log.exception("reportlab failed to build the Excel PDF")
        raise ToolError("We couldn't convert this workbook because of an unsupported layout.", 422)
    return buf.getvalue()


# ----------------------------------------------------------------------------
# Images -> PDF
# ----------------------------------------------------------------------------


def _load_image(data: bytes, number: int) -> Image.Image:
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("error", Image.DecompressionBombWarning)
            img = Image.open(io.BytesIO(data))
            img.load()
    except (Image.DecompressionBombError, Image.DecompressionBombWarning):
        raise ToolError(f"Image {number} has too many pixels to process.", 413)
    except Exception:
        raise ToolError(f"Image {number} couldn't be read. Use JPG or PNG files.", 400)

    img = ImageOps.exif_transpose(img)  # respect the phone's rotation flag
    if img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info):
        rgba = img.convert("RGBA")
        background = Image.new("RGB", rgba.size, (255, 255, 255))
        background.paste(rgba, mask=rgba.getchannel("A"))
        img = background
    else:
        img = img.convert("RGB")
    img.thumbnail((4200, 4200))
    return img


def images_to_pdf(files: list[bytes]) -> bytes:
    if not files:
        raise ToolError("Upload at least one image.", 400)
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf, pageCompression=1)
    margin = 24
    for number, data in enumerate(files, start=1):
        img = _load_image(data, number)
        iw, ih = img.size
        page = A4 if ih >= iw else landscape(A4)
        pw, ph = page
        scale = min((pw - 2 * margin) / iw, (ph - 2 * margin) / ih)
        w, h = iw * scale, ih * scale
        jpg = io.BytesIO()
        img.save(jpg, "JPEG", quality=90, optimize=True)  # JPEG streams are embedded as-is
        jpg.seek(0)
        pdf.setPageSize(page)
        pdf.drawImage(ImageReader(jpg), (pw - w) / 2, (ph - h) / 2, w, h)
        pdf.showPage()
    pdf.save()
    return buf.getvalue()


# ----------------------------------------------------------------------------
# Merge / split / compress
# ----------------------------------------------------------------------------


def merge_pdfs(files: list[bytes]) -> bytes:
    if len(files) < 2:
        raise ToolError("Upload at least 2 PDFs to merge.", 400)
    writer = pypdf.PdfWriter()
    for number, data in enumerate(files, start=1):
        try:
            writer.append(open_pdf(data))
        except ToolError as exc:
            raise ToolError(f"File {number}: {exc.message}", exc.status)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def parse_ranges(spec: str, total: int) -> list[list[int]]:
    """'1-3, 5, 8-10' -> [[0,1,2],[4],[7,8,9]] (zero-based page indexes)."""
    spec = re.sub(r"\s*[-\u2013\u2014]\s*", "-", spec.strip())
    groups: list[list[int]] = []
    for part in re.split(r"[,\u060c;\s]+", spec):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not m:
            raise ToolError(f'"{part}" is not a valid page range. Use a format like 1-3,5,8-10.', 400)
        first = int(m.group(1))
        last = int(m.group(2) or first)
        if first < 1 or last < first or last > total:
            raise ToolError(f"Page range {part} is outside this PDF (it has {total} pages).", 400)
        groups.append(list(range(first - 1, last)))
    return groups


def split_pdf(content: bytes, ranges: str = "") -> bytes:
    """One PDF per page, or one PDF per requested range, returned as a ZIP."""
    reader = open_pdf(content)
    total = len(reader.pages)
    groups = parse_ranges(ranges, total) if ranges.strip() else [[i] for i in range(total)]
    if not groups:
        raise ToolError("No pages selected.", 400)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for number, pages in enumerate(groups, start=1):
            writer = pypdf.PdfWriter()
            for index in pages:
                writer.add_page(reader.pages[index])
            part = io.BytesIO()
            writer.write(part)
            if ranges.strip():
                label = f"{pages[0] + 1}" if len(pages) == 1 else f"{pages[0] + 1}-{pages[-1] + 1}"
                name = f"part_{number:02d}_pages_{label}.pdf"
            else:
                name = f"page_{number:03d}.pdf"
            zf.writestr(name, part.getvalue())
    return buf.getvalue()


# level -> (downscale factor, JPEG quality, only resize images larger than this many px)
_REDUCE_LEVELS = {
    "low": (0.85, 85, 1600),
    "medium": (0.6, 65, 1000),
    "high": (0.45, 45, 800),
}


def reduce_pdf(content: bytes, level: str = "medium") -> tuple[bytes, int, int]:
    """Recompress images and streams. Returns (pdf_bytes, original_size, result_size).

    If the result is not smaller, the original bytes are returned unchanged.
    """
    scale, quality, threshold = _REDUCE_LEVELS.get(level, _REDUCE_LEVELS["medium"])
    reader = open_pdf(content)
    writer = pypdf.PdfWriter(clone_from=reader)

    for page in writer.pages:
        try:
            images = list(page.images)
        except Exception as exc:
            log.debug("Could not list images on a page: %s", exc)
            images = []
        for image_file in images:
            try:
                ref = getattr(image_file, "indirect_reference", None)
                obj = ref.get_object() if ref is not None else {}
                if "/SMask" in obj or "/Mask" in obj:  # replacing would drop transparency
                    continue
                pil = image_file.image
                if pil.mode == "CMYK":
                    pil = pil.convert("RGB")
                if pil.mode not in ("RGB", "L"):
                    continue
                width, height = pil.size
                if width * height < 40_000:
                    continue
                if max(width, height) > threshold:
                    pil = pil.resize(
                        (max(1, int(width * scale)), max(1, int(height * scale))),
                        Image.Resampling.LANCZOS,
                    )
                image_file.replace(pil, quality=quality)
            except Exception as exc:
                log.debug("Skipping an image that could not be recompressed: %s", exc)
        try:
            page.compress_content_streams()
        except Exception as exc:
            log.debug("Could not compress a content stream: %s", exc)

    if hasattr(writer, "compress_identical_objects"):
        try:
            writer.compress_identical_objects(remove_identicals=True, remove_orphans=True)
        except Exception as exc:
            log.debug("compress_identical_objects failed: %s", exc)

    out = io.BytesIO()
    writer.write(out)
    result = out.getvalue()
    if len(result) >= len(content):
        return content, len(content), len(content)
    return result, len(content), len(result)